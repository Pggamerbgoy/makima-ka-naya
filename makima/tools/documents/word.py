"""
Makima v2 — Document Processing: Word Document Engine.
Architectural Pattern: Hermes-style modular document tools with workspace boundary clamping.
Zero-Hardcoding Guarantee: Fully dynamic path resolution via Makima filesystem boundary resolver.
Library-first: Direct use of python-docx for rich Word document generation and structured Markdown extraction.
"""

from __future__ import annotations

import logging
from typing import Any, Dict, List

import docx

from makima.tools.filesystem.safe_io import get_workspace_dir, resolve_safe_path
from makima.tools.registry import tool

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Registered Tools
# ---------------------------------------------------------------------------

@tool(
    name="create_word_document",
    description="Creates a rich Word document (.docx) with structured sections, headings, formatted paragraphs, bullet lists, and styled tables.",
    tags=["documents", "word", "office"],
)
def create_word_document(
    file_path: str,
    title: str,
    sections: List[Dict[str, Any]],
) -> Dict[str, Any]:
    """
    Creates a Word document (.docx) populated with structured sections.

    Args:
        file_path: Destination path relative to the workspace boundary (e.g. 'docs/specification.docx').
        title: Document title (styled as Title / Level 0 heading).
        sections: List of section dicts. Each section can contain:
                  - 'heading': Section title string.
                  - 'level': Heading depth (1 to 4, defaults to 1).
                  - 'text': Main paragraph text string.
                  - 'paragraphs': Optional list of multiple text paragraph strings.
                  - 'bullets': Optional list of bullet point strings.
                  - 'table': Optional dict with 'headers' (list[str]) and 'rows' (list[list[Any]]).

    Returns:
        Dictionary containing document creation status, file metadata, and structure counts.

    Raises:
        BoundaryTraversalError: If file_path resolves outside workspace boundary.
        ValueError: If title is empty.
    """
    if not title:
        raise ValueError("Cannot create Word document with an empty title.")

    resolved = resolve_safe_path(file_path)
    if resolved.suffix.lower() != ".docx":
        resolved = resolved.with_suffix(".docx")

    # Ensure target parent directories exist
    resolved.parent.mkdir(parents=True, exist_ok=True)

    doc = docx.Document()

    # 1. Document Title
    doc.add_heading(title, level=0)

    # 2. Iterate through sections
    for sec in sections:
        if not isinstance(sec, dict):
            continue

        # Heading
        heading = sec.get("heading")
        if heading:
            level = max(1, min(int(sec.get("level", 1)), 4))
            doc.add_heading(str(heading), level=level)

        # Single paragraph text
        text = sec.get("text")
        if text:
            doc.add_paragraph(str(text))

        # Multiple paragraphs
        paragraphs = sec.get("paragraphs")
        if isinstance(paragraphs, list):
            for p in paragraphs:
                if p:
                    doc.add_paragraph(str(p))

        # Bullet list
        bullets = sec.get("bullets")
        if isinstance(bullets, list):
            for b in bullets:
                if b is not None:
                    doc.add_paragraph(str(b), style="List Bullet")

        # Table
        table_spec = sec.get("table")
        if isinstance(table_spec, dict):
            headers = table_spec.get("headers", [])
            rows = table_spec.get("rows", [])

            if headers:
                cols_count = len(headers)
                t = doc.add_table(rows=len(rows) + 1, cols=cols_count)
                t.style = "Table Grid"

                # Style header row
                header_row = t.rows[0]
                for col_idx, h in enumerate(headers):
                    cell = header_row.cells[col_idx]
                    cell.text = str(h)
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.bold = True

                # Fill data rows
                for row_idx, row_values in enumerate(rows):
                    curr_row = t.rows[row_idx + 1]
                    for col_idx in range(cols_count):
                        val = row_values[col_idx] if col_idx < len(row_values) else ""
                        curr_row.cells[col_idx].text = str(val) if val is not None else ""

                # Spacing paragraph after table
                doc.add_paragraph()

    # Save to disk
    doc.save(resolved)

    ws_root = get_workspace_dir().resolve()
    rel_path = str(resolved.relative_to(ws_root)).replace("\\", "/")

    return {
        "status": "success",
        "file_path": rel_path,
        "absolute_path": str(resolved),
        "title": title,
        "sections_count": len(sections),
        "paragraphs_count": len(doc.paragraphs),
        "tables_count": len(doc.tables),
        "file_size_bytes": resolved.stat().st_size,
    }


@tool(
    name="read_word_document",
    description="Reads a Word document (.docx) into structured paragraphs, tables, and a unified Markdown representation.",
    tags=["documents", "word", "office"],
)
def read_word_document(file_path: str) -> Dict[str, Any]:
    """
    Reads paragraphs, headings, bullet lists, and tables from a Word document into structured text.

    Args:
        file_path: Path to the Word document relative to the workspace boundary.

    Returns:
        Dictionary containing extracted paragraphs, extracted tables, and synthesized Markdown text.

    Raises:
        BoundaryTraversalError: If file_path resolves outside workspace boundary.
        FileNotFoundError: If the target file does not exist.
    """
    resolved = resolve_safe_path(file_path)

    if not resolved.exists():
        raise FileNotFoundError(f"Word document not found: '{file_path}' (resolved: '{resolved}')")

    if resolved.is_dir():
        raise IsADirectoryError(f"Target path is a directory, not a file: '{file_path}'")

    doc = docx.Document(resolved)

    extracted_paragraphs: List[Dict[str, str]] = []
    markdown_lines: List[str] = []

    # 1. Process Paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue

        style_name = p.style.name if p.style else "Normal"
        extracted_paragraphs.append({
            "text": txt,
            "style": style_name,
        })

        if style_name == "Title":
            markdown_lines.append(f"# {txt}\n")
        elif style_name.startswith("Heading 1"):
            markdown_lines.append(f"## {txt}\n")
        elif style_name.startswith("Heading 2"):
            markdown_lines.append(f"### {txt}\n")
        elif style_name.startswith("Heading 3"):
            markdown_lines.append(f"#### {txt}\n")
        elif "Bullet" in style_name:
            markdown_lines.append(f"- {txt}")
        else:
            markdown_lines.append(f"{txt}\n")

    # 2. Process Tables
    extracted_tables: List[Dict[str, Any]] = []

    for t_idx, table in enumerate(doc.tables):
        if not table.rows:
            continue

        headers = [cell.text.strip() for cell in table.rows[0].cells]
        rows: List[List[str]] = []

        for row in table.rows[1:]:
            row_vals = [cell.text.strip() for cell in row.cells]
            rows.append(row_vals)

        extracted_tables.append({
            "table_index": t_idx + 1,
            "headers": headers,
            "rows": rows,
            "rows_count": len(rows),
            "columns_count": len(headers),
        })

        # Render Markdown Table
        if headers:
            markdown_lines.append("\n" + "| " + " | ".join(headers) + " |")
            markdown_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for r in rows:
                markdown_lines.append("| " + " | ".join(r) + " |")
            markdown_lines.append("")

    unified_markdown = "\n".join(markdown_lines).strip()
    ws_root = get_workspace_dir().resolve()
    rel_path = str(resolved.relative_to(ws_root)).replace("\\", "/")

    return {
        "status": "success",
        "file_path": rel_path,
        "paragraphs": extracted_paragraphs,
        "tables": extracted_tables,
        "markdown": unified_markdown,
        "paragraphs_count": len(doc.paragraphs),
        "tables_count": len(doc.tables),
    }
