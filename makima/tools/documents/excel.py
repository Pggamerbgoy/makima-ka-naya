"""
Makima v2 — Document Processing: Excel Spreadsheet Engine.
Architectural Pattern: Hermes-style modular document tools with workspace boundary clamping.
Zero-Hardcoding Guarantee: Fully dynamic path resolution via Makima filesystem boundary resolver.
Library-first: Direct use of openpyxl for high-fidelity styled Excel generation and reading.
"""

from __future__ import annotations

import logging
from datetime import date, datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from makima.tools.filesystem.safe_io import get_workspace_dir, resolve_safe_path
from makima.tools.registry import tool

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Excel Styling Constants
# ---------------------------------------------------------------------------

HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

ROW_FONT = Font(name="Segoe UI", size=10)
ZEBRA_FILL = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

THIN_BORDER_SIDE = Side(style="thin", color="D9D9D9")
CELL_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE,
)


# ---------------------------------------------------------------------------
# Registered Tools
# ---------------------------------------------------------------------------

@tool(
    name="create_excel_sheet",
    description="Creates a styled Excel workbook (.xlsx) with custom headers, styled data rows, auto-adjusted column widths, and frozen header panes.",
    tags=["documents", "excel", "office"],
)
def create_excel_sheet(
    file_path: str,
    headers: List[str],
    rows: List[List[Any]],
    sheet_title: str = "Sheet1",
) -> Dict[str, Any]:
    """
    Creates an Excel workbook (.xlsx) with styled headers and data rows.

    Args:
        file_path: Destination path relative to the workspace boundary (e.g. 'reports/sales.xlsx').
        headers: List of column header strings.
        rows: List of data rows, where each row is a list of values.
        sheet_title: Title of the primary sheet (defaults to 'Sheet1').

    Returns:
        Dictionary containing creation status, file metadata, and row/column counts.

    Raises:
        BoundaryTraversalError: If file_path resolves outside workspace boundary.
        ValueError: If headers list is empty.
    """
    if not headers:
        raise ValueError("Cannot create Excel sheet with empty headers list.")

    resolved = resolve_safe_path(file_path)
    if resolved.suffix.lower() != ".xlsx":
        resolved = resolved.with_suffix(".xlsx")

    # Ensure target parent directories exist
    resolved.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title or "Sheet1"

    # Ensure grid lines are visible
    if ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True

    # 1. Write Header Row
    ws.append(headers)
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = CELL_BORDER

    # 2. Write Data Rows with Formatting
    for row_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)
        current_fill = ZEBRA_FILL if is_even else WHITE_FILL

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = ROW_FONT
            cell.fill = current_fill
            cell.border = CELL_BORDER

            val = cell.value
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif isinstance(val, (datetime, date)):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "YYYY-MM-DD"
            elif isinstance(val, bool):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # 3. Freeze Header Row
    ws.freeze_panes = "A2"

    # 4. Auto-adjust Column Widths
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(headers[col_idx - 1] or ""))

        for row_data in rows:
            if col_idx - 1 < len(row_data):
                val = row_data[col_idx - 1]
                val_len = len(str(val)) if val is not None else 0
                if val_len > max_len:
                    max_len = val_len

        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save to disk
    wb.save(resolved)

    ws_root = get_workspace_dir().resolve()
    rel_path = str(resolved.relative_to(ws_root)).replace("\\", "/")

    return {
        "status": "success",
        "file_path": rel_path,
        "absolute_path": str(resolved),
        "sheet_title": ws.title,
        "rows_count": len(rows),
        "columns_count": len(headers),
        "file_size_bytes": resolved.stat().st_size,
    }


@tool(
    name="read_excel_sheet",
    description="Reads data from an Excel spreadsheet into a structured dictionary of headers, rows, and sheet names.",
    tags=["documents", "excel", "office"],
)
def read_excel_sheet(
    file_path: str,
    sheet_title: Optional[str] = None,
    max_rows: int = 100,
) -> Dict[str, Any]:
    """
    Reads rows from an Excel file into a structured list of dictionaries.

    Args:
        file_path: Path to the target Excel file relative to the workspace boundary.
        sheet_title: Specific sheet name to read (defaults to active sheet if omitted).
        max_rows: Maximum number of data rows to retrieve (defaults to 100).

    Returns:
        Dictionary containing sheet headers, structured row records, and sheet metadata.

    Raises:
        BoundaryTraversalError: If file_path resolves outside workspace boundary.
        FileNotFoundError: If the target file does not exist.
        ValueError: If the specified sheet_title does not exist in workbook.
    """
    resolved = resolve_safe_path(file_path)

    if not resolved.exists():
        raise FileNotFoundError(f"Excel file not found: '{file_path}' (resolved: '{resolved}')")

    if resolved.is_dir():
        raise IsADirectoryError(f"Target path is a directory, not a file: '{file_path}'")

    wb = openpyxl.load_workbook(resolved, data_only=True)

    if sheet_title:
        if sheet_title not in wb.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_title}' not found in workbook '{file_path}'. "
                f"Available sheets: {wb.sheetnames}"
            )
        ws = wb[sheet_title]
    else:
        ws = wb.active

    # Extract headers from row 1
    raw_headers = [cell.value for cell in ws[1]] if ws.max_row >= 1 else []
    headers = [str(h) if h is not None else f"Column_{i+1}" for i, h in enumerate(raw_headers)]

    rows_data: List[Dict[str, Any]] = []

    # Read up to max_rows data rows
    for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, max_rows + 1), values_only=True):
        # Skip completely empty rows
        if all(v is None for v in row):
            continue

        row_dict: Dict[str, Any] = {}
        for idx, header in enumerate(headers):
            val = row[idx] if idx < len(row) else None
            # Format datetime / date objects cleanly
            if isinstance(val, (datetime, date)):
                val = val.isoformat()
            row_dict[header] = val
        rows_data.append(row_dict)

    ws_root = get_workspace_dir().resolve()
    rel_path = str(resolved.relative_to(ws_root)).replace("\\", "/")

    return {
        "status": "success",
        "file_path": rel_path,
        "sheet_title": ws.title,
        "sheet_names": wb.sheetnames,
        "headers": headers,
        "rows": rows_data,
        "total_rows_read": len(rows_data),
        "total_sheet_rows": max(0, ws.max_row - 1) if ws.max_row else 0,
        "total_sheet_columns": ws.max_column or len(headers),
    }
