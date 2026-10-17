"""
Makima v2 — Document Processing: PowerPoint Presentation Engine.
Architectural Pattern: Hermes-style modular document tools with workspace boundary clamping.
Zero-Hardcoding Guarantee: Fully dynamic path resolution via Makima filesystem boundary resolver.
Library-first: Direct use of python-pptx for modern 16:9 widescreen PowerPoint deck generation and inspection.
"""

from __future__ import annotations

import logging
from typing import Any, Dict, List, Optional

import pptx
from pptx.util import Inches

from makima.tools.filesystem.safe_io import get_workspace_dir, resolve_safe_path
from makima.tools.registry import tool

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Registered Tools
# ---------------------------------------------------------------------------

@tool(
    name="create_powerpoint",
    description="Creates a modern 16:9 widescreen PowerPoint presentation (.pptx) with title slide, content slides, formatted bullet points, and speaker notes.",
    tags=["documents", "powerpoint", "presentation", "office"],
)
def create_powerpoint(
    file_path: str,
    title: str,
    subtitle: str = "",
    slides: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, Any]:
    """
    Creates a PowerPoint presentation (.pptx) with title slide and content slides.

    Args:
        file_path: Destination path relative to the workspace boundary (e.g. 'presentations/pitch.pptx').
        title: Main presentation title displayed on the title slide.
        subtitle: Optional subtitle text for the title slide.
        slides: Optional list of content slide dictionaries. Each dict can contain:
                - 'title': Slide heading string.
                - 'content': Optional main body text.
                - 'bullets': Optional list of bullet point strings.
                - 'notes': Optional presenter/speaker notes.

    Returns:
        Dictionary containing presentation metadata, slide counts, and file size.

    Raises:
        BoundaryTraversalError: If file_path resolves outside workspace boundary.
        ValueError: If title is empty.
    """
    if not title:
        raise ValueError("Cannot create PowerPoint presentation with an empty title.")

    resolved = resolve_safe_path(file_path)
    if resolved.suffix.lower() != ".pptx":
        resolved = resolved.with_suffix(".pptx")

    # Ensure target parent directories exist
    resolved.parent.mkdir(parents=True, exist_ok=True)

    prs = pptx.Presentation()

    # Set modern 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Slide 0: Title Slide
    title_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(title_layout)

    title_shape = title_slide.shapes.title
    title_shape.text = title

    if subtitle and len(title_slide.placeholders) > 1:
        subtitle_shape = title_slide.placeholders[1]
        subtitle_shape.text = subtitle

    # 2. Slides 1..N: Content Slides
    content_slides_data = slides or []
    content_layout = prs.slide_layouts[1]

    for s_data in content_slides_data:
        if not isinstance(s_data, dict):
            continue

        slide = prs.slides.add_slide(content_layout)

        # Set Slide Title
        slide_title = s_data.get("title", "")
        if slide.shapes.title:
            slide.shapes.title.text = str(slide_title)

        # Populate Content Body / Bullets
        body_shape = slide.placeholders[1] if len(slide.placeholders) > 1 else None

        bullets = s_data.get("bullets")
        content_text = s_data.get("content")

        if body_shape:
            tf = body_shape.text_frame
            tf.word_wrap = True

            if isinstance(bullets, list) and bullets:
                # First bullet replaces default initial paragraph
                tf.paragraphs[0].text = str(bullets[0])
                tf.paragraphs[0].level = 0

                for b in bullets[1:]:
                    p = tf.add_paragraph()
                    p.text = str(b)
                    p.level = 0
            elif content_text:
                tf.paragraphs[0].text = str(content_text)

        # Add Speaker Notes if provided
        notes_text = s_data.get("notes")
        if notes_text:
            notes_slide = slide.notes_slide
            tf_notes = notes_slide.notes_text_frame
            tf_notes.text = str(notes_text)

    # Save to disk
    prs.save(resolved)

    ws_root = get_workspace_dir().resolve()
    rel_path = str(resolved.relative_to(ws_root)).replace("\\", "/")

    return {
        "status": "success",
        "file_path": rel_path,
        "absolute_path": str(resolved),
        "title": title,
        "subtitle": subtitle,
        "slides_count": len(prs.slides),
        "file_size_bytes": resolved.stat().st_size,
    }


@tool(
    name="read_powerpoint",
    description="Reads a PowerPoint presentation (.pptx) into a structured summary of slides, titles, bullet contents, and speaker notes.",
    tags=["documents", "powerpoint", "presentation", "office"],
)
def read_powerpoint(file_path: str) -> Dict[str, Any]:
    """
    Reads slides from a PowerPoint (.pptx) file into a structured dictionary.

    Args:
        file_path: Path to the PowerPoint presentation relative to the workspace boundary.

    Returns:
        Dictionary containing presentation slides, text frames, and speaker notes.

    Raises:
        BoundaryTraversalError: If file_path resolves outside workspace boundary.
        FileNotFoundError: If the target file does not exist.
    """
    resolved = resolve_safe_path(file_path)

    if not resolved.exists():
        raise FileNotFoundError(f"PowerPoint file not found: '{file_path}' (resolved: '{resolved}')")

    if resolved.is_dir():
        raise IsADirectoryError(f"Target path is a directory, not a file: '{file_path}'")

    prs = pptx.Presentation(resolved)

    slides_data: List[Dict[str, Any]] = []

    for idx, slide in enumerate(prs.slides):
        slide_title = slide.shapes.title.text.strip() if slide.shapes.title else ""
        content_lines: List[str] = []

        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = p.text.strip()
                    if txt:
                        content_lines.append(txt)

        # Extract speaker notes
        notes_text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()

        slides_data.append({
            "slide_number": idx + 1,
            "title": slide_title,
            "contents": content_lines,
            "notes": notes_text,
        })

    ws_root = get_workspace_dir().resolve()
    rel_path = str(resolved.relative_to(ws_root)).replace("\\", "/")

    return {
        "status": "success",
        "file_path": rel_path,
        "slides_count": len(prs.slides),
        "slides": slides_data,
    }
